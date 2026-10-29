from __future__ import annotations

import io

from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from bordereaux.mapping.service import generate_mapping_candidates
from bordereaux.cleaning.service import acknowledge_cleaning, run_cleaning
from bordereaux.core.exceptions import NotFoundError, StateError
from bordereaux.financial_checks.service import proceed_to_report_generation, run_financial_checks
from bordereaux.ingestion.service import create_processing_run_with_upload, validate_processing_run
from bordereaux.mapping.service import approve_candidate, finalize_mappings, list_candidates
from bordereaux.models.project import Project, User
from bordereaux.reporting.docx_report import DOCX_CONTENT_TYPE
from bordereaux.reporting.service import generate_reports, resolve_download
from bordereaux.storage.local import LocalStorageService
from bordereaux.treaty.schemas import TreatyFiguresRequest
from bordereaux.treaty.service import submit_treaty_figures

pytestmark = pytest.mark.unit

from tests.unit.financial_checks.test_financial_checks_service import _CSV  # reuse the full-coverage fixture


async def _run_to_report_generation_pending(db_session: AsyncSession, storage: LocalStorageService, project: Project, user: User):
    result = await create_processing_run_with_upload(db_session, storage, project.id, "f.csv", _CSV, user.id)
    run = result["run"]
    await validate_processing_run(db_session, storage, run.id, user.id)
    await run_cleaning(db_session, storage, run.id, user.id)
    await acknowledge_cleaning(db_session, storage, run.id, user.id)
    await generate_mapping_candidates(db_session, storage, run.id, user.id)
    for c in await list_candidates(db_session, run.id):
        if c.status == "PROPOSED":
            await approve_candidate(db_session, c.id, user.id)
    await finalize_mappings(db_session, storage, run.id, user.id)
    await submit_treaty_figures(
        db_session, run.id, TreatyFiguresRequest(commission_pct="0.22", tax_pct="0.11", cession_pct="0.40"), user.id
    )
    await run_financial_checks(db_session, storage, run.id, user.id)
    await proceed_to_report_generation(db_session, run.id, user.id)
    return run


async def test_generate_reports_produces_real_artifacts(
    db_session: AsyncSession, storage: LocalStorageService, project: Project, user: User
):
    run = await _run_to_report_generation_pending(db_session, storage, project, user)
    processed, audit = await generate_reports(db_session, storage, run.id, user.id)
    assert processed.format == "CSV"
    assert audit.format == "JSON"

    content, filename, content_type = await resolve_download(db_session, storage, run.id, "processed_file")
    assert filename == "processed_file.csv"
    assert b"gross_premium" in content

    content, filename, content_type = await resolve_download(db_session, storage, run.id, "audit_report")
    import json

    payload = json.loads(content)
    assert payload["final_status"] == "COMPLETED"
    assert payload["financial_checks"]["total_findings"] >= 1
    assert payload["source"]["filename"] == "f.csv"

    content, filename, content_type = await resolve_download(db_session, storage, run.id, "audit_report_human")
    assert filename == "audit_report.docx"
    assert content_type == DOCX_CONTENT_TYPE
    document = Document(io.BytesIO(content))
    text = " ".join(p.text for p in document.paragraphs)
    cells = " ".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert "Audit Report" in text and "Summary" in text
    assert "f.csv" in cells  # source file appears in the summary table
    assert "Gross Premium" in cells or "gross_premium" in cells  # mapping table
    assert len(document.tables) >= 5


async def test_generate_reports_requires_correct_state(
    db_session: AsyncSession, storage: LocalStorageService, project: Project, user: User
):
    result = await create_processing_run_with_upload(db_session, storage, project.id, "f.csv", _CSV, user.id)
    with pytest.raises(StateError):
        await generate_reports(db_session, storage, result["run"].id, user.id)


async def test_download_unknown_run_is_not_found(db_session: AsyncSession, storage: LocalStorageService):
    with pytest.raises(NotFoundError):
        await resolve_download(db_session, storage, uuid4(), "processed_file")


async def test_download_before_generation_is_not_found(
    db_session: AsyncSession, storage: LocalStorageService, project: Project, user: User
):
    run = await _run_to_report_generation_pending(db_session, storage, project, user)
    with pytest.raises(NotFoundError):
        await resolve_download(db_session, storage, run.id, "processed_file")
