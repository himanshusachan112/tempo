"""The Word audit report: every section becomes a table, and confidence + evidence are shown."""

from __future__ import annotations

import io

import pytest
from docx import Document

from bordereaux.reporting.docx_report import build_audit_docx, describe_evidence

pytestmark = pytest.mark.unit

PAYLOAD = {
    "run_id": "r1", "project_id": "p1", "generated_at": "2026-10-28T10:00:00+00:00", "final_status": "COMPLETED",
    "source": {"filename": "bdx.xlsx", "checksum": "abc", "format": "XLSX", "size_bytes": 2048},
    "cleaning": {"attempts": 1, "latest_attempt": 1, "before_row_count": 10, "after_row_count": 9,
                 "before_column_count": 7, "after_column_count": 7,
                 "transformations": [{"type": "trim_whitespace", "column": "Insured", "count": 3}], "warnings": []},
    "mapping": {"schema_id": "s", "finalized_at": None, "coverage_required": "1.0", "coverage_optional": "0.5",
                "approved_mappings": [{
                    "source_column": "Gross Premium", "target_field": "gross_premium", "origin": "EXACT",
                    "confidence_total": "0.9500",
                    "confidence_breakdown": {"semantic": 0.95, "value": 1, "context": 0.8, "layout": 0.5, "ai": 0.5},
                    "evidence": [{"type": "alias_match", "stage": "EXACT", "score": 0.95}]}]},
    "treaty_figures": {"commission_pct": "0.25", "tax_pct": "0", "cession_pct": "0.4", "fx_rates": [{"currency": "EUR", "rate": "1.1"}]},
    "money_relations": [{"label": "Net", "expression": "net = gross - comm", "origin": "USER", "rows_evaluated": 9, "rows_failed": 1}],
    "column_renamings": [{"rule_id": "MCR-D1", "source_column": "Gross Premium", "original_meaning": "gross_premium",
                          "new_meaning": "ceded_premium", "evidence": {"text": "ceded share", "agreement": "1.0"}}],
    "financial_checks": {"total_findings": 1, "status_counts": {"UNRESOLVED": 1}, "category_counts": {"TREATY_RATIO": 1}},
    "discrepancies": [{"check_id": "MCR-B1", "category": "TREATY_RATIO", "severity": "ERROR", "status": "UNRESOLVED",
                       "affected_row_number": 4, "affected_field": "commission", "expected_value": "250", "observed_value": "300",
                       "variance": "50", "message": "Commission should be 250", "evidence": {"row_agreement": "0.9"}, "override": None}],
    "ai_activity": [{"purpose": "MAPPING", "status": "OK", "model": "m", "items_sent": 3, "items_resolved": 2, "latency_ms": 120, "message": "ok"}],
    "audit_trail": [{"timestamp": "2026-10-28T09:00:00+00:00", "event_type": "reports.generated", "action": "CREATE", "actor_id": None, "reason": None}],
}


def _text(doc: Document) -> tuple[str, str]:
    return (
        "\n".join(p.text for p in doc.paragraphs),
        " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells),
    )


def test_every_section_is_present_with_confidence_and_evidence() -> None:
    doc = Document(io.BytesIO(build_audit_docx(PAYLOAD)))
    text, cells = _text(doc)
    for heading in ("Summary", "Data cleaning", "Column mapping, confidence and evidence", "Treaty figures applied",
                    "Money-column relationships", "Financial checks and discrepancies", "AI assistance", "Audit trail"):
        assert heading in text
    assert "95% (High)" in cells and "Alias match (Stage EXACT, Score 0.95)" in cells  # confidence + evidence
    assert "Commission should be 250" in cells and "Row agreement 0.9" in cells
    assert "ceded share" in cells and "EUR" in cells


def test_a_sparse_payload_still_renders_without_empty_sections() -> None:
    doc = Document(io.BytesIO(build_audit_docx({"run_id": "r", "final_status": "COMPLETED"})))
    text, _ = _text(doc)
    assert "Summary" in text and "AI assistance" not in text and "Treaty figures applied" not in text


def test_control_characters_in_data_cannot_break_the_file() -> None:
    payload = {**PAYLOAD, "source": {"filename": "bad\x00name\x0b.csv"}}
    assert Document(io.BytesIO(build_audit_docx(payload)))


def test_describe_evidence() -> None:
    assert describe_evidence(None) == "None recorded"
    assert describe_evidence({"type": "manual_override", "description": "x"}) == "Manual override"
