"""Suggested column renamings through the real HTTP API (rule doc §5.1, MCR-D1).

The contract the UI relies on: suggestions appear once the discrepancy review is open, accepting one
produces a *separate* revised list, and the original discrepancy findings never change.
"""

from __future__ import annotations

from pathlib import Path

import httpx
import openpyxl
import pytest

from tests.conftest import auth
from tests.integration import _pipeline as pipeline
from tests.integration.test_ai_and_relations_api import finish_mapping, start, use_model

pytestmark = pytest.mark.integration

TOKEN = pipeline.TOKEN


@pytest.fixture
def source(tmp_path: Path) -> Path:
    """'Gross Premium' is really the 40% ceded share of a 10,000 gross; commission is 25% of it (§5.1 worked example)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bordereaux"
    ws.append(["UMR", "Insured", "Sum Insured", "Gross Premium", "Commission", "Net Premium"])
    for i in range(1, 9):
        ws.append([f"B{i}2024XY000{i}", f"Insured {i}", 1_000_000, 4000, 1000, 3000])
    path = tmp_path / "ceded_source.xlsx"
    wb.save(path)
    return path


async def reach_review(client: httpx.AsyncClient, project, source: Path, monkeypatch) -> str:
    use_model(monkeypatch, None)
    base = await start(client, project, source)
    await client.post(f"{base}/mappings/generate", params={"use_ai": "false"}, headers=auth(TOKEN))
    await finish_mapping(client, base)
    body = {
        "commission_pct": "0.25", "tax_pct": "0", "cession_pct": "0.4", "expected_si_rate": "0.01",
        "settlement_currency": "USD", "fx_rates": [],
    }
    assert (await client.post(f"{base}/treaty-figures", json=body, headers=auth(TOKEN))).status_code == 201
    assert (await client.post(f"{base}/financial-checks/run", headers=auth(TOKEN))).status_code == 200
    return base


async def test_suggestion_accept_revise_and_revert_leave_the_original_findings_untouched(client, project, source, monkeypatch) -> None:
    base = await reach_review(client, project, source, monkeypatch)
    before = (await client.get(f"{base}/discrepancies", headers=auth(TOKEN))).json()["data"]["discrepancies"]

    ws = (await client.get(f"{base}/column-renamings", headers=auth(TOKEN))).json()["data"]
    d1 = next(s for s in ws["suggestions"] if s["rule_id"] == "MCR-D1")
    assert d1["field"] == "gross_premium" and d1["suggested_role"] == "ceded_premium" and d1["renamable"]
    assert d1["flag"] == "MEANING_CONFLICT" and not d1["accepted"]
    assert "ceded share, not the full gross" in d1["evidence_text"]
    assert ws["editable"] and not ws["revised"]["active"]

    r = await client.post(f"{base}/column-renamings/accept", json={"keys": [d1["key"]], "note": "cedant confirmed"}, headers=auth(TOKEN))
    assert r.status_code == 200, r.text
    revised = r.json()["data"]
    assert revised["accepted_count"] == 1 and revised["revised"]["active"]
    assert not any(f["check_id"] == "MCR-D1" for f in revised["revised"]["findings"])
    assert revised["revised"]["summary"]["original_total"] == len(before)
    assert any(s["accepted"] for s in revised["suggestions"])

    after = (await client.get(f"{base}/discrepancies", headers=auth(TOKEN))).json()["data"]["discrepancies"]
    assert after == before  # the original table is exactly as the financial checks left it

    again = await client.post(f"{base}/column-renamings/accept", json={"keys": [d1["key"]]}, headers=auth(TOKEN))
    assert again.status_code == 200 and again.json()["data"]["accepted_count"] == 1  # idempotent

    r = await client.delete(f"{base}/column-renamings", params={"key": d1["key"]}, headers=auth(TOKEN))
    assert r.status_code == 200 and not r.json()["data"]["revised"]["active"]


async def test_unknown_or_missing_keys_are_rejected(client, project, source, monkeypatch) -> None:
    base = await reach_review(client, project, source, monkeypatch)
    bad = await client.post(f"{base}/column-renamings/accept", json={"keys": ["MCR-D1:nope"]}, headers=auth(TOKEN))
    assert bad.status_code == 400
    empty = await client.post(f"{base}/column-renamings/accept", json={"keys": []}, headers=auth(TOKEN))
    assert empty.status_code == 400
    gone = await client.delete(f"{base}/column-renamings", params={"key": "MCR-D1:gross_premium"}, headers=auth(TOKEN))
    assert gone.status_code == 404
